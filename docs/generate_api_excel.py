import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

navy = '002A54'
white = 'FFFFFF'
header_font = Font(name='Calibri', bold=True, color=white, size=11)
header_fill = PatternFill(start_color=navy, end_color=navy, fill_type='solid')
data_font = Font(name='Calibri', size=10)
border = Border(
    left=Side(style='thin', color='D0D5DD'),
    right=Side(style='thin', color='D0D5DD'),
    top=Side(style='thin', color='D0D5DD'),
    bottom=Side(style='thin', color='D0D5DD')
)
wrap = Alignment(wrap_text=True, vertical='center')
center = Alignment(horizontal='center', vertical='center')

method_colors = {
    'GET': ('1E40AF', 'DBEAFE'),
    'POST': ('166534', 'DCFCE7'),
    'PUT': ('854D0E', 'FEF9C3'),
    'PATCH': ('6B21A8', 'F3E8FF'),
    'DELETE': ('991B1B', 'FEE2E2'),
}

# Sheet 1: API Endpoints
ws = wb.active
ws.title = 'API Endpoints'
ws.sheet_properties.tabColor = navy

headers = ['#', 'Category', 'HTTP Method', 'Endpoint', 'File', 'Purpose']
widths = [5, 20, 12, 55, 35, 40]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    col_letter = openpyxl.utils.get_column_letter(i)
    ws.column_dimensions[col_letter].width = w
    cell = ws.cell(row=1, column=i, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

endpoints = [
    ('Authentication', 'POST', '/api/v1/tenant/company-auth/login', 'AuthApiService.cs', 'User login with email/password'),
    ('Authentication', 'POST', '/api/v1/tenant/device/auth/register-device', 'AuthApiService.cs', 'Register device with license key'),
    ('Authentication', 'POST', '/api/v1/tenant/device/auth/validate-device', 'AuthApiService.cs', 'Validate registered device'),
    ('Authentication', 'POST', '/api/v1/tenant/device/auth/refresh', 'AuthApiService.cs', 'Refresh device access token'),
    ('Authentication', 'POST', '/api/v1/tenant/device/auth/admin-refresh', 'AuthApiService.cs', 'Refresh admin access token'),
    ('Authentication', 'POST', '/api/v1/tenant/device/auth/admin-login', 'AuthApiService.cs', 'Admin login on registered device'),
    ('Sessions', 'POST', '/api/v1/Revit/session/Open', 'SessionSyncService.cs', 'Create/open new session'),
    ('Sessions', 'PATCH', '/api/v1/Revit/session/{sessionId}', 'SessionSyncService.cs', 'Update existing session'),
    ('Sessions', 'PATCH', '/api/v1/Revit/session/{sessionId}/heartbeat', 'SessionSyncService.cs', 'Send heartbeat with CPU/RAM/GPU'),
    ('Sessions', 'PATCH', '/api/v1/Revit/session/{sessionId}/role-swap', 'SessionSyncService.cs', 'Swap user role in session'),
    ('Sessions', 'GET', '/api/v1/Revit/session', 'SessionSyncService.cs', 'Fetch all sessions'),
    ('Sessions', 'GET', '/api/v1/Revit/session/{sessionId}', 'SessionSyncService.cs', 'Fetch specific session'),
    ('Sessions', 'PATCH', '/api/v1/Revit/unmonitored-users', 'SessionSyncService.cs', 'Update unmonitored users list'),
    ('Models', 'POST', '/api/v1/Revit/models/register', 'ModelSyncService.cs', 'Register/sync model'),
    ('Models', 'GET', '/api/v1/Revit/models', 'ModelSyncService.cs', 'Fetch all registered models'),
    ('Models', 'GET', '/api/v1/Revit/models/{modelGuid}', 'ModelSyncService.cs', 'Fetch single model'),
    ('Models', 'PATCH', '/api/v1/Revit/models/{modelGuid}/deactivate', 'ModelSyncService.cs', 'Deactivate model'),
    ('Models', 'PATCH', '/api/v1/Revit/models/{modelGuid}/activate', 'ModelSyncService.cs', 'Activate model'),
    ('Models', 'GET', '/api/v1/Revit/models/StaticModelInfo/{modelGuid}', 'ModelHealthDashboard.xaml.cs', 'Combined metrics + thresholds'),
    ('Model Sessions', 'POST', '/api/v1/Revit/models/model-sessions', 'ModelSessionSyncService.cs', 'Create model session'),
    ('Model Sessions', 'GET', '/api/v1/Revit/models/model-sessions', 'ModelSessionSyncService.cs', 'Fetch all model sessions'),
    ('Model Sessions', 'GET', '/api/v1/Revit/models/{modelGuid}/sessions', 'ModelSessionSyncService.cs', 'Fetch sessions for model'),
    ('Model Sessions', 'PATCH', '/api/v1/Revit/models/model-sessions/{sessionId}', 'ModelSessionSyncService.cs', 'Update model session'),
    ('Metrics', 'POST', '/api/v1/Revit/metrics/syncsave', 'MetricsSyncService.cs', 'Upload sync/save metrics'),
    ('Metrics', 'POST', '/api/v1/Revit/metrics/periodic', 'MetricsSyncService.cs', 'Upload periodic metrics'),
    ('Metrics', 'POST', '/api/v1/Revit/metrics/manual', 'MetricsSyncService.cs', 'Upload manual metrics'),
    ('Metrics', 'POST', '/api/v1/Revit/model-syncs', 'MetricsSyncService.cs', 'Record sync event'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/syncsave/latest', 'MetricsSyncService.cs', 'Latest sync/save by model'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/periodic/latest', 'MetricsSyncService.cs', 'Latest periodic by model'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/manual/latest', 'MetricsSyncService.cs', 'Latest manual by model'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/syncsave', 'MetricsSyncService.cs', 'Fetch all sync/save'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/periodic', 'MetricsSyncService.cs', 'Fetch all periodic'),
    ('Metrics', 'GET', '/api/v1/Revit/metrics/manual', 'MetricsSyncService.cs', 'Fetch all manual'),
    ('Command Protection', 'POST', '/api/v1/Revit/command-protections', 'CommandProtectionSyncService.cs', 'Create command protection'),
    ('Command Protection', 'PUT', '/api/v1/Revit/command-protections/{id}', 'CommandProtectionSyncService.cs', 'Update command protection'),
    ('Command Protection', 'DELETE', '/api/v1/Revit/command-protections/{id}', 'CommandProtectionSyncService.cs', 'Delete command protection'),
    ('Command Protection', 'GET', '/api/v1/Revit/command-protections', 'CommandProtectionSyncService.cs', 'Fetch all command protections'),
    ('Command Protection', 'GET', '/api/v1/Revit/command-protections/by-model/{modelGuid}', 'CommandProtectionSyncService.cs', 'Fetch by model'),
    ('Event Protection', 'GET', '/api/v1/Revit/event-protections/by-model/{modelGuid}', 'EventProtectionSyncService.cs', 'Fetch event protections for model'),
    ('Event Protection', 'POST', '/api/v1/Revit/event-protections', 'EventProtectionSyncService.cs', 'Create event protection'),
    ('Event Protection', 'PUT', '/api/v1/Revit/event-protections/{id}', 'EventProtectionSyncService.cs', 'Update event protection'),
    ('Event Protection', 'PATCH', '/api/v1/Revit/event-protections/{id}/is-enabled', 'EventProtectionSyncService.cs', 'Toggle enabled state'),
    ('Pin Protection', 'GET', '/api/v1/Revit/pin-protections/by-model/{modelGuid}', 'PinProtectionSyncService.cs', 'Fetch pins for model'),
    ('Pin Protection', 'GET', '/api/v1/Revit/pin-protections/{id}', 'PinProtectionSyncService.cs', 'Fetch specific pin'),
    ('Pin Protection', 'POST', '/api/v1/Revit/pin-protections', 'PinProtectionSyncService.cs', 'Create pin protection'),
    ('Pin Protection', 'DELETE', '/api/v1/Revit/pin-protections/{id}', 'PinProtectionSyncService.cs', 'Delete pin protection'),
    ('Rule Protection', 'POST', '/api/v1/Revit/rule-protections', 'RulesSyncService.cs', 'Create rule'),
    ('Rule Protection', 'PUT', '/api/v1/Revit/rule-protections/{id}', 'RulesSyncService.cs', 'Update rule'),
    ('Rule Protection', 'DELETE', '/api/v1/Revit/rule-protections/{id}', 'RulesSyncService.cs', 'Delete rule'),
    ('Rule Protection', 'GET', '/api/v1/Revit/rule-protections/by-model/{modelGuid}', 'RulesSyncService.cs', 'Fetch rules for model'),
    ('Health Monitor', 'GET', '/api/v1/Revit/health-monitor-protections/by-model/{modelGuid}', 'HealthMonitorProtectionSyncService.cs', 'Model thresholds'),
    ('Audit & Evidence', 'POST', '/api/v1/Revit/audit-logs', 'AuditLogSyncService.cs', 'Upload audit entries'),
    ('Audit & Evidence', 'POST', '/api/v1/Revit/audit-logs/{id}/send-mail', 'AuditLogMailDispatchService.cs', 'Trigger email notification'),
    ('Audit & Evidence', 'POST', '/api/v1/Revit/evidence-images/with-images', 'EvidenceUploadQueue.cs', 'Upload screenshots (multipart)'),
    ('OTP', 'POST', '/api/v1/Revit/otp/generate', 'OtpGenerateDialog.xaml.cs', 'Generate one-time password'),
    ('OTP', 'POST', '/api/v1/Revit/otp/validate', 'OtpRepository.cs', 'Validate OTP code'),
    ('Revit Data', 'GET', '/api/v1/Revit/commands', 'RevitApiService.cs', 'Fetch all Revit commands'),
    ('Revit Data', 'GET', '/api/v1/Revit/commands/{name}', 'RevitApiService.cs', 'Fetch command details'),
    ('Revit Data', 'GET', '/api/v1/Revit/categories', 'RevitApiService.cs', 'Fetch all categories'),
    ('Revit Data', 'GET', '/api/v1/Revit/categories/codes?categoryName={name}', 'RevitApiService.cs', 'Fetch category code'),
    ('Revit Data', 'GET', '/api/v1/Revit/categories/codes?names={names}', 'RevitApiService.cs', 'Batch category codes'),
    ('Admin', 'POST', '/api/v1/Revit/project-model-admins/model', 'ModelAdminsSyncService.cs', 'Sync model admin relations'),
    ('AI', 'GET', '/api/v1/master/ai-training', 'ApiKnowledgeProvider.cs', 'Fetch AI training data'),
    ('Support', 'POST', '/api/v1/master/tickets/revit', 'ReportIssueDialog.xaml.cs', 'Submit bug report (multipart)'),
]

row = 2
for i, (cat, method, endpoint, file, purpose) in enumerate(endpoints, 1):
    ws.cell(row=row, column=1, value=i).font = data_font
    ws.cell(row=row, column=1).alignment = center
    ws.cell(row=row, column=1).border = border

    ws.cell(row=row, column=2, value=cat).font = data_font
    ws.cell(row=row, column=2).border = border

    mc = ws.cell(row=row, column=3, value=method)
    fg, bg = method_colors.get(method, ('000000', 'FFFFFF'))
    mc.font = Font(name='Calibri', bold=True, size=10, color=fg)
    mc.fill = PatternFill(start_color=bg, end_color=bg, fill_type='solid')
    mc.alignment = center
    mc.border = border

    ws.cell(row=row, column=4, value=endpoint).font = Font(name='Consolas', size=10)
    ws.cell(row=row, column=4).border = border

    ws.cell(row=row, column=5, value=file).font = Font(name='Calibri', size=10, color='475569')
    ws.cell(row=row, column=5).border = border

    ws.cell(row=row, column=6, value=purpose).font = data_font
    ws.cell(row=row, column=6).alignment = wrap
    ws.cell(row=row, column=6).border = border

    row += 1

# Sheet 2: SignalR
ws2 = wb.create_sheet('SignalR Hub Methods')
ws2.sheet_properties.tabColor = '166534'
headers2 = ['#', 'Hub Method', 'File', 'Purpose']
widths2 = [5, 30, 30, 45]
for i, (h, w) in enumerate(zip(headers2, widths2), 1):
    col_letter = openpyxl.utils.get_column_letter(i)
    ws2.column_dimensions[col_letter].width = w
    cell = ws2.cell(row=1, column=i, value=h)
    cell.font = Font(name='Calibri', bold=True, color=white, size=11)
    cell.fill = PatternFill(start_color='166534', end_color='166534', fill_type='solid')
    cell.alignment = center
    cell.border = border

signalr = [
    ('SendSessionData', 'SignalRService.cs', 'Session activity updates'),
    ('SendModelData', 'SignalRService.cs', 'Model registration updates'),
    ('SendManualMetrics', 'SignalRService.cs', 'Manual metrics notifications'),
    ('SendPeriodicMetrics', 'SignalRService.cs', 'Periodic metrics notifications'),
    ('SendSyncSaveMetrics', 'SignalRService.cs', 'Sync/save metrics notifications'),
    ('SendRuleData', 'SignalRService.cs', 'Rule protection updates'),
    ('SendModelAdminData', 'SignalRService.cs', 'Model admin updates'),
    ('SendModelSessionData', 'SignalRService.cs', 'Model session updates'),
]

for i, (method, file, purpose) in enumerate(signalr, 1):
    ws2.cell(row=i+1, column=1, value=i).font = data_font
    ws2.cell(row=i+1, column=1).alignment = center
    ws2.cell(row=i+1, column=1).border = border
    ws2.cell(row=i+1, column=2, value=method).font = Font(name='Consolas', size=10, bold=True)
    ws2.cell(row=i+1, column=2).border = border
    ws2.cell(row=i+1, column=3, value=file).font = Font(name='Calibri', size=10, color='475569')
    ws2.cell(row=i+1, column=3).border = border
    ws2.cell(row=i+1, column=4, value=purpose).font = data_font
    ws2.cell(row=i+1, column=4).border = border

# Sheet 3: Service Files
ws3 = wb.create_sheet('Service Files')
ws3.sheet_properties.tabColor = '854D0E'
headers3 = ['#', 'File', 'Path', 'Purpose']
widths3 = [5, 40, 30, 35]
for i, (h, w) in enumerate(zip(headers3, widths3), 1):
    col_letter = openpyxl.utils.get_column_letter(i)
    ws3.column_dimensions[col_letter].width = w
    cell = ws3.cell(row=1, column=i, value=h)
    cell.font = Font(name='Calibri', bold=True, color=white, size=11)
    cell.fill = PatternFill(start_color='854D0E', end_color='854D0E', fill_type='solid')
    cell.alignment = center
    cell.border = border

files = [
    ('AuthApiService.cs', 'Infrastructure/Auth/', 'Authentication'),
    ('AuthenticatedHttpClient.cs', 'Infrastructure/Auth/', 'HTTP client with JWT'),
    ('SessionSyncService.cs', 'Infrastructure/Api/', 'Sessions'),
    ('ModelSyncService.cs', 'Infrastructure/Api/', 'Models'),
    ('ModelSessionSyncService.cs', 'Infrastructure/Api/', 'Model sessions'),
    ('MetricsSyncService.cs', 'Infrastructure/Api/', 'Metrics'),
    ('CommandProtectionSyncService.cs', 'Infrastructure/Api/', 'Command protection'),
    ('EventProtectionSyncService.cs', 'Infrastructure/Api/', 'Event protection'),
    ('PinProtectionSyncService.cs', 'Infrastructure/Api/', 'Pin protection'),
    ('RulesSyncService.cs', 'Infrastructure/Api/', 'Rule protection'),
    ('HealthMonitorProtectionSyncService.cs', 'Infrastructure/Api/', 'Health thresholds'),
    ('AuditLogSyncService.cs', 'Infrastructure/Api/', 'Audit logs'),
    ('AuditLogMailDispatchService.cs', 'Infrastructure/Api/', 'Audit email'),
    ('EvidenceUploadQueue.cs', 'Core/Evidence/', 'Screenshot upload'),
    ('RevitApiService.cs', 'Infrastructure/Api/', 'Revit commands/categories'),
    ('ModelAdminsSyncService.cs', 'Infrastructure/Api/', 'Model admins'),
    ('ApiKnowledgeProvider.cs', 'AI/Knowledge/', 'AI training data'),
    ('SignalRService.cs', 'Infrastructure/SignalR/Core/', 'Real-time hub'),
]

for i, (f, path, purpose) in enumerate(files, 1):
    ws3.cell(row=i+1, column=1, value=i).font = data_font
    ws3.cell(row=i+1, column=1).alignment = center
    ws3.cell(row=i+1, column=1).border = border
    ws3.cell(row=i+1, column=2, value=f).font = Font(name='Consolas', size=10)
    ws3.cell(row=i+1, column=2).border = border
    ws3.cell(row=i+1, column=3, value=path).font = Font(name='Calibri', size=10, color='475569')
    ws3.cell(row=i+1, column=3).border = border
    ws3.cell(row=i+1, column=4, value=purpose).font = data_font
    ws3.cell(row=i+1, column=4).border = border

ws.freeze_panes = 'A2'
ws2.freeze_panes = 'A2'
ws3.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:F{row-1}'

wb.save('docs/ZeManage_API_Endpoints.xlsx')
print('Excel created: docs/ZeManage_API_Endpoints.xlsx')
