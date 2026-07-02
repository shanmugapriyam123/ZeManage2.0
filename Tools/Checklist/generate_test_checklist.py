"""
ZeManage Revit Add-in — Test Checklist Generator (v2 — Accurate / Button-mapped)
Produces: ZeManage_Test_Checklist.xlsx
Modules map 1-to-1 with ZeManage ribbon panels.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from collections import Counter

# ── Colour palette ────────────────────────────────────────────────────────────
HEADER_BG  = "1E293B"
MOD_COLORS = {
    "Protection":       ("1D4ED8", "DBEAFE"),
    "Activity Tracker": ("065F46", "D1FAE5"),
    "Health Monitor":   ("B45309", "FEF3C7"),
    "Sync Control":     ("6D28D9", "EDE9FE"),
    "AI":               ("0E7490", "CFFAFE"),
    "Addons":           ("BE123C", "FFE4E6"),
    "Support":          ("374151", "F1F5F9"),
}

# ── Style helpers ─────────────────────────────────────────────────────────────
def thin(): s = Side(style="thin", color="CBD5E1"); return Border(left=s, right=s, top=s, bottom=s)
def med():  s = Side(style="medium", color="94A3B8"); return Border(left=s, right=s, top=s, bottom=s)
def fill(c): return PatternFill("solid", fgColor=c)
def fnt(bold=False, color="1E293B", size=9): return Font(bold=bold, color=color, size=size, name="Calibri")

# ── COLUMNS ───────────────────────────────────────────────────────────────────
COLUMNS = [
    ("A", "#",               5),
    ("B", "Module",         18),
    ("C", "Sub-Module",     22),
    ("D", "Test Case ID",   14),
    ("E", "Test Description", 38),
    ("F", "Pre-conditions / Steps", 52),
    ("G", "Expected Result",        50),
    ("H", "Priority",       11),
    ("I", "Model Type",     14),   # Local / Cloud / Both / N/A
    ("J", "Status",         12),
    ("K", "Actual Result / Notes",  44),
    ("L", "Tester",         14),
    ("M", "Date Tested",    13),
    ("N", "Build",          13),
]

# ── TEST DATA ─────────────────────────────────────────────────────────────────
# (Module, Sub-Module, ID, Description, Pre-conditions/Steps, Expected Result, Priority, Model-Type)
TESTS = [

    # ═══════════════════════════════════════════════════════════════════════════
    # 1. PROTECTION  (Ribbon Panel: "Protection")
    #    Buttons: Pin Protection | Command Protection | Event Restriction | Rule Management
    # ═══════════════════════════════════════════════════════════════════════════

    # ── 1a. Pin Protection ──────────────────────────────────────────────────
    ("Protection", "Pin Protection", "PP-01",
     "Normal pin — element pinned without protection mode",
     "1. Open workshared model\n2. Select any element\n3. ZeManage → Pin Protection → Standard Pin",
     "Element pinned (pushpin icon visible). No protection mode applied. No OTP required to unpin.",
     "High", "Both"),

    ("Protection", "Pin Protection", "PP-02",
     "Protected pin — Notify mode applied",
     "1. Select element → Pin Protection → Protected Pin\n2. Set mode = Notify → Save",
     "Element pinned with Notify badge. Unpin attempt triggers audit log entry. No OTP.",
     "High", "Both"),

    ("Protection", "Pin Protection", "PP-03",
     "Protected pin — Prevent mode requires OTP to unpin",
     "1. Pin element with Prevent mode\n2. Select pinned element → Unpin",
     "OTP dialog appears. Correct OTP → unpin proceeds. Wrong OTP → unpin blocked. Audit log created.",
     "High", "Both"),

    ("Protection", "Pin Protection", "PP-04",
     "Unpin with RequireCommentForUnpin = true",
     "1. Pin element with Prevent + RequireComment enabled\n2. Enter correct OTP",
     "Comment field shown after OTP validation. Unpin only completes after comment entered.",
     "Medium", "Both"),

    ("Protection", "Pin Protection", "PP-05",
     "Evidence screenshots captured on protected unpin",
     "1. Unpin a Prevent-mode pinned element\n2. Check %TEMP%\\BIManage\\Evidence\\{sessionId}\\",
     "Before + After PNG files saved. AuditLogId linked in evidence_captures table.",
     "High", "Both"),

    ("Protection", "Pin Protection", "PP-06",
     "Pin record POSTed to server API",
     "1. Create a protected pin\n2. Check server / log for API call",
     "POST /api/v1/Revit/pin-protections sent. Payload has pinProtectionId, elementGuid, protectedBy, modelGuid.",
     "High", "Both"),

    ("Protection", "Pin Protection", "PP-07",
     "Deactivated pin synced to server on unpin",
     "1. Unpin a server-synced protected element\n2. Check API logs",
     "PATCH /api/v1/Revit/pin-protections/{id} with isActive=false. Server reflects deactivation.",
     "High", "Both"),

    # ── 1b. Command Protection ──────────────────────────────────────────────
    ("Protection", "Command Protection", "CP-01",
     "Command Protection dialog opens (admin only)",
     "1. Log in as Admin\n2. ZeManage → Command Protection",
     "Command Protection dialog opens with list of commands and their current modes.",
     "High", "Both"),

    ("Protection", "Command Protection", "CP-02",
     "Monitor mode — command executes with audit log",
     "1. Set Delete command to Monitor mode\n2. Log in as regular user → select element → Delete",
     "Delete executes. Audit log entry created with mode=Notify, user, element info.",
     "High", "Both"),

    ("Protection", "Command Protection", "CP-03",
     "Guide mode — intervention dialog shown before execution",
     "1. Set Move to Guide mode\n2. Regular user → select element → Move",
     "Intervention dialog appears. User can Proceed or Cancel. Audit log records choice.",
     "High", "Both"),

    ("Protection", "Command Protection", "CP-04",
     "Prevent mode — command fully blocked",
     "1. Set Copy to Prevent mode\n2. Regular user → select element → Copy",
     "Command blocked. Dialog explains reason. No copy made. Audit log records block.",
     "High", "Both"),

    ("Protection", "Command Protection", "CP-05",
     "Admin bypasses Prevent mode",
     "1. Command set to Prevent\n2. Log in as Admin → attempt same command",
     "Admin can proceed without block. Non-admin is still blocked.",
     "High", "Both"),

    ("Protection", "Command Protection", "CP-06",
     "Bulk command settings sync from server (live update)",
     "1. Admin changes protection mode on server/web\n2. Wait one SignalR push or sync cycle",
     "Client reflects updated mode within one cycle. No Revit restart required.",
     "High", "Both"),

    # ── 1c. Event Restriction ────────────────────────────────────────────────
    ("Protection", "Event Restriction", "ER-01",
     "Event Restriction dialog opens (admin only)",
     "1. Log in as Admin\n2. ZeManage → Event Restriction",
     "Dialog opens listing Revit events with their current protection modes.",
     "High", "Both"),

    ("Protection", "Event Restriction", "ER-02",
     "View Activation event — Monitor mode logs view switches",
     "1. Set View Activated event to Monitor\n2. Switch between views in a project",
     "Audit log entry created for each view activation: view name, user, timestamp.",
     "High", "Both"),

    ("Protection", "Event Restriction", "ER-03",
     "View Activation event — Prevent mode blocks view switch",
     "1. Set View Activated event to Prevent\n2. Attempt to switch to a restricted view",
     "View switch blocked. User remains on current view. Audit log records block.",
     "Medium", "Both"),

    ("Protection", "Event Restriction", "ER-04",
     "Family Load event — Guide mode shows dialog before load",
     "1. Set Family Loading to Guide\n2. Load a family into the project",
     "Intervention dialog appears before family loads. Cancel → load aborted. Proceed → family loads.",
     "Medium", "Both"),

    ("Protection", "Event Restriction", "ER-05",
     "Document Saving event — Monitor logs every save",
     "1. Set DocumentSaving to Monitor\n2. Ctrl+S (local save)",
     "Audit log entry created with document name, user, timestamp.",
     "Medium", "Both"),

    ("Protection", "Event Restriction", "ER-06",
     "Document Closing event — Prevent blocks document close",
     "1. Set DocumentClosing to Prevent\n2. Attempt to close the document",
     "Close blocked. Document remains open. Audit log records attempt.",
     "Medium", "Both"),

    # ── 1d. Rule Management ──────────────────────────────────────────────────
    ("Protection", "Rule Management", "RM-01",
     "Rule Management dialog opens and loads existing rules",
     "1. Log in as Admin\n2. ZeManage → Rule Management",
     "Dialog opens. Existing rules listed with name, category filter, mode, status.",
     "High", "Both"),

    ("Protection", "Rule Management", "RM-02",
     "Create rule scoped to element category — Delete protected for Walls",
     "1. Create rule: category=Walls, command=Delete, mode=Prevent\n2. Select a Wall → Delete",
     "Delete blocked for Walls. Non-wall elements not affected. Audit log records block.",
     "High", "Both"),

    ("Protection", "Rule Management", "RM-03",
     "Create rule scoped to specific Family Type",
     "1. Create rule: family type = 'Structural Column 300x300', command=Move, mode=Guide\n2. Select matching column → Move",
     "Guide dialog appears for matching column. Non-matching types execute freely.",
     "High", "Both"),

    ("Protection", "Rule Management", "RM-04",
     "Rule priority — higher priority rule wins on conflict",
     "1. Create two overlapping rules for same element (different modes, different priorities)\n2. Execute command",
     "Rule with higher priority value wins. Result matches expected priority order.",
     "Medium", "Both"),

    ("Protection", "Rule Management", "RM-05",
     "Disabled rule is skipped — no protection applied",
     "1. Create a rule → disable it via toggle\n2. Trigger the same command",
     "No intervention. Disabled rule fully bypassed in evaluation pipeline.",
     "Medium", "Both"),

    ("Protection", "Rule Management", "RM-06",
     "Rules filter — search and status filter work correctly",
     "1. Open Rule Management\n2. Type rule name in search box\n3. Filter by Active/Inactive",
     "Search narrows results in real time. Status filter shows only matching rules.",
     "Low", "Both"),

    ("Protection", "Rule Management", "RM-07",
     "Rule sync from server — live update without restart",
     "1. Admin modifies a rule on server\n2. Wait for SignalR push or manual sync cycle",
     "Updated rule reflected in protection engine within one cycle. No Revit restart.",
     "High", "Both"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 2. ACTIVITY TRACKER  (Ribbon Panel: "Activity Tracker")
    #    Buttons: Session Info | Crash Register | Model Activities
    # ═══════════════════════════════════════════════════════════════════════════

    # ── 2a. Session Info ─────────────────────────────────────────────────────
    ("Activity Tracker", "Session Info", "SI-01",
     "Session Info dialog shows current session details",
     "1. ZeManage → Session Info",
     "Dialog shows: session ID, user, machine, Revit version, login time, active documents.",
     "High", "N/A"),

    ("Activity Tracker", "Session Info", "SI-02",
     "Model opening time is recorded and displayed",
     "1. Close all docs → open a workshared model\n2. Open Session Info → check document entry",
     "Model opening time (opened_at) and opening duration (ms) shown for the active document.",
     "High", "Both"),

    ("Activity Tracker", "Session Info", "SI-03",
     "Session synced to server on startup",
     "1. Launch Revit + ZeManage\n2. Check log for session POST",
     "POST /api/v1/Revit/session with sessionId, machineId, revitUsername. Server returns 200.",
     "High", "N/A"),

    ("Activity Tracker", "Session Info", "SI-04",
     "Document session opened on model open",
     "1. Open a workshared model\n2. Check document_sessions table in bimanage.db",
     "Row inserted: sessionId FK, modelGuid, modelName, openedAt, status=Active.",
     "High", "Both"),

    ("Activity Tracker", "Session Info", "SI-05",
     "Document session closed on model close",
     "1. Close a workshared model\n2. Check document_sessions table",
     "closedAt populated. Status = Closed. Duration minutes calculated correctly.",
     "High", "Both"),

    # ── 2b. Crash Register ───────────────────────────────────────────────────
    ("Activity Tracker", "Crash Register", "CR-01",
     "Crash Register dialog opens and lists crash history",
     "1. ZeManage → Crash Register",
     "Dialog opens. Previous crash entries listed with timestamp, session ID, model name.",
     "High", "N/A"),

    ("Activity Tracker", "Crash Register", "CR-02",
     "Session crash detected after force-kill",
     "1. Open Revit + ZeManage\n2. Kill Revit via Task Manager\n3. Re-launch Revit",
     "On next startup: previous session flagged is_crashed=1 in document_sessions. Crash Register shows entry.",
     "High", "N/A"),

    ("Activity Tracker", "Crash Register", "CR-03",
     "Model session crash detection (model-level)",
     "1. Open a workshared model\n2. Force-kill Revit\n3. Re-launch + open Crash Register",
     "Model document session also marked as crashed. Crash Register distinguishes session crash from model crash.",
     "High", "Both"),

    ("Activity Tracker", "Crash Register", "CR-04",
     "Crash notification sent to server",
     "1. Force-kill and re-launch Revit\n2. Check API logs",
     "PATCH /api/v1/Revit/session/{id} with crashDetected=1 sent on startup recovery.",
     "High", "N/A"),

    ("Activity Tracker", "Crash Register", "CR-05",
     "Normal close does NOT register as crash",
     "1. Close Revit normally (File → Exit)\n2. Re-launch → open Crash Register",
     "Previous session: is_crashed=0, status=Closed. No crash entry appears in Crash Register.",
     "High", "N/A"),

    # ── 2c. Model Activities ─────────────────────────────────────────────────
    ("Activity Tracker", "Model Activities", "MA-01",
     "Model Activities shows current users on same model",
     "1. Have two users open same workshared model\n2. User A: ZeManage → Model Activities",
     "Both users listed as active on the model. Revit usernames and machine IDs shown.",
     "High", "Both"),

    ("Activity Tracker", "Model Activities", "MA-02",
     "User leaving model shown in real time (SignalR)",
     "1. User B closes the model while Model Activities dialog is open\n2. Observe dialog",
     "User B's entry updates to Closed/Inactive in real time via SignalR without reopening dialog.",
     "High", "Both"),

    ("Activity Tracker", "Model Activities", "MA-03",
     "Sync history visible in Model Activities",
     "1. Perform 3 syncs on a model\n2. Open Model Activities",
     "All 3 sync records listed with user, start time, duration, status (Completed/Failed).",
     "High", "Both"),

    ("Activity Tracker", "Model Activities", "MA-04",
     "Model opening time shown in activity list",
     "1. Open model\n2. Check Model Activities for opening entry",
     "Model open event visible with timestamp matching Session Info opened_at value.",
     "Medium", "Both"),

    ("Activity Tracker", "Model Activities", "MA-05",
     "Model type (local vs cloud) shown correctly",
     "1. Open a local workshared model → check Model Activities\n2. Repeat with cloud (BIM360/ACC) model",
     "Local: path shows UNC/drive path. Cloud: shows cloud project name (e.g. BIM 360://ProjectName).",
     "Medium", "Both"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 3. HEALTH MONITOR  (Ribbon Panel: "Health Monitor")
    #    Buttons: Analyze Model | Model Health
    # ═══════════════════════════════════════════════════════════════════════════

    # ── 3a. Model Health Dashboard ───────────────────────────────────────────
    ("Health Monitor", "Model Health Dashboard", "MH-01",
     "Model Health dashboard opens for current model",
     "1. Open any project model\n2. ZeManage → Model Health",
     "Dashboard opens showing model name, GUID, and health score ring with colour coding.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-02",
     "Health score calculated from 5 critical parameters",
     "1. Open Model Health for a model with known metric values\n2. Verify score",
     "Score = (passed parameters / 5) × 100. Checks: Warnings, Duplicates, In-Place Families, Views Not on Sheets, File Size.\n≥75% = Green, 50-74% = Orange, <50% = Red.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-03",
     "Warnings count — Pass vs Needs Attention classification",
     "1. Model with >100 warnings\n2. Compare against configured goal threshold",
     "If warnings > MaximumWarningCount goal: 'Needs Attention'. If ≤ goal: 'Pass'.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-04",
     "Duplicate elements — classification against threshold",
     "1. Check model with known duplicate element count (from warnings analysis)\n2. Open dashboard",
     "DuplicateElementsCount compared against MaxDuplicateElementsCount goal. Correct Pass/Needs Attention.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-05",
     "In-Place Family count — classification against threshold",
     "1. Model with known in-place families\n2. Check dashboard",
     "InplaceFamiliesCount compared against MaxInPlaceFamilyCount (default: 10). Correct classification.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-06",
     "Views Not On Sheets — correct count and classification",
     "1. Create views not placed on sheets\n2. Sync model → open dashboard",
     "ViewsNotOnSheetsCount reflects actual count. Compared against ViewsNotOnSheet goal.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-07",
     "File size — LOCAL model shows actual file size",
     "1. Open a local workshared model (UNC path)\n2. Open dashboard → check File Size",
     "File size in MB matches actual file size on disk (within ±5%). Not zero.",
     "High", "Local"),

    ("Health Monitor", "Model Health Dashboard", "MH-08",
     "File size — CLOUD model uses CollaborationCache fallback",
     "1. Open a cloud model (BIM 360 or ACC)\n2. Open dashboard → check File Size",
     "File size shown from local CollaborationCache copy. If cache not found, 0 shown (acceptable).",
     "High", "Cloud"),

    ("Health Monitor", "Model Health Dashboard", "MH-09",
     "Goals/thresholds loaded from server API",
     "1. Admin sets custom thresholds on server (MaxWarnings=50, MaxInPlace=5)\n2. Open dashboard",
     "Dashboard uses server-configured thresholds. Pass/Needs Attention recalculates accordingly.",
     "High", "Both"),

    ("Health Monitor", "Model Health Dashboard", "MH-10",
     "Metrics history displayed in dashboard",
     "1. Perform 3+ syncs (auto-captures fast metrics)\n2. Open Model Health",
     "Historical metric records shown in timeline/table. Most recent first. Timestamps accurate.",
     "Medium", "Both"),

    # ── 3b. Sync/Save Metrics (Fast) — Auto-captured ────────────────────────
    ("Health Monitor", "Sync/Save Metrics (Auto)", "HM-01",
     "Fast metrics captured automatically on every sync",
     "1. Sync a workshared model\n2. Check model_file_metrics_sync_save table",
     "Row inserted with captureType='sync', fileSizeBytes, warningsCount, totalViewsCount, etc. capturedAt matches sync time.",
     "High", "Both"),

    ("Health Monitor", "Sync/Save Metrics (Auto)", "HM-02",
     "Fast metrics captured on local save",
     "1. Ctrl+S (local save, not sync)\n2. Check model_file_metrics_sync_save",
     "Row inserted with captureType='save'. File size and other fast metrics recorded.",
     "High", "Local"),

    ("Health Monitor", "Sync/Save Metrics (Auto)", "HM-03",
     "Sync/Save metrics POSTed to API",
     "1. Sync model\n2. Check log for metrics API call",
     "POST /api/v1/Revit/metrics/syncsave sent. Returns 200. Offline queued if server unavailable.",
     "High", "Both"),

    ("Health Monitor", "Sync/Save Metrics (Auto)", "HM-04",
     "File size = 0 for cloud model sync metrics (acceptable if no cache)",
     "1. Sync a cloud model\n2. Check fast metrics row → fileSizeBytes",
     "fileSizeBytes = 0 if CollaborationCache not found. Non-zero if cache hit. Both are acceptable.",
     "Medium", "Cloud"),

    ("Health Monitor", "Sync/Save Metrics (Auto)", "HM-05",
     "Shared coordinates captured in sync metrics",
     "1. Open model with shared coordinates\n2. Sync → check metrics row",
     "SharedCoordNs, SharedCoordEw, SharedCoordElevation populated. Not scientific notation in API payload.",
     "Medium", "Both"),

    # ── 3c. Periodic Metrics (Medium) — Auto daily ───────────────────────────
    ("Health Monitor", "Periodic Metrics (Daily)", "PM-01",
     "Periodic metrics auto-captured on daily schedule (2 AM UTC)",
     "1. Let Revit run overnight\n2. Check model_file_metrics_periodic next morning",
     "Row inserted around 02:00 UTC with is_manual_trigger=0. 10 metrics populated: totalElements, views not on sheets, etc.",
     "High", "Both"),

    ("Health Monitor", "Periodic Metrics (Daily)", "PM-02",
     "Manual periodic capture via CapturePeriodicMetrics (developer feature)",
     "1. Developer panel → CapturePeriodicMetrics button\n2. Confirm dialog → check results",
     "10 metrics collected immediately. is_manual_trigger=1. POST to /api/v1/Revit/metrics/periodic.",
     "Medium", "Both"),

    ("Health Monitor", "Periodic Metrics (Daily)", "PM-03",
     "Unplaced and unenclosed rooms detected correctly",
     "1. Model with known unplaced/unenclosed rooms\n2. Trigger periodic capture",
     "UnplacedRoomsCount and UnenclosedRoomsCount match known counts in model.",
     "Medium", "Both"),

    ("Health Monitor", "Periodic Metrics (Daily)", "PM-04",
     "Views not on sheets count is accurate",
     "1. Count views not placed on sheets manually\n2. Compare with periodic metric",
     "ViewsNotOnSheetsCount matches manual count. Includes 3D/section/elevation views not on sheets.",
     "Medium", "Both"),

    # ── 3d. Manual/Expensive Metrics ────────────────────────────────────────
    ("Health Monitor", "Analyze Model (Manual)", "AM-01",
     "Analyze Model confirmation dialog shown before analysis",
     "1. ZeManage → Analyze Model",
     "Confirmation dialog lists what will be analyzed. User must click Analyze to proceed.",
     "High", "Both"),

    ("Health Monitor", "Analyze Model (Manual)", "AM-02",
     "Progress dialog shows 0–100% during analysis",
     "1. Confirm analysis\n2. Observe progress dialog",
     "Progress advances: 0-40% (families), 40-95% (purgeable check), 95-100% finalize. No freeze >30s.",
     "High", "Both"),

    ("Health Monitor", "Analyze Model (Manual)", "AM-03",
     "Families >5MB count (proxy: >1000 instances) captured correctly",
     "1. Model with known heavy family (>1000 instances)\n2. Run analysis\n3. Check results",
     "FamiliesOver5MbCount reflects families with >1000 instances. Zero if none exist.",
     "High", "Both"),

    ("Health Monitor", "Analyze Model (Manual)", "AM-04",
     "Purgeable elements count is accurate",
     "1. Model with known unused types (family symbols, wall types, fill patterns)\n2. Run analysis",
     "PurgeableElementsCount matches count of unused elements across 9 type categories.",
     "High", "Both"),

    ("Health Monitor", "Analyze Model (Manual)", "AM-05",
     "Analysis results saved to DB and synced to API",
     "1. Run analysis\n2. Check model_file_metrics_manual table and API logs",
     "Row in model_file_metrics_manual. POST /api/v1/Revit/metrics/manual sent. commandSource='ribbon'.",
     "High", "Both"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 4. SYNC CONTROL  (Ribbon Panel: "Sync Control")
    #    Buttons: Sync Settings | Sync Queue
    # ═══════════════════════════════════════════════════════════════════════════

    # ── 4a. Sync Settings ────────────────────────────────────────────────────
    ("Sync Control", "Sync Settings", "SS-01",
     "Sync Settings dialog opens and loads current settings from DB",
     "1. ZeManage → Sync Settings",
     "Dialog opens with current values: enabled toggle, mode, sync interval, idle timeout, schedule.",
     "High", "N/A"),

    ("Sync Control", "Sync Settings", "SS-02",
     "Background sync enable/disable toggle persists",
     "1. Toggle Enable Background Sync off → Save\n2. Reopen dialog",
     "Toggle persists as disabled. Engine stops firing. Toggle on → engine resumes.",
     "High", "N/A"),

    ("Sync Control", "Sync Settings", "SS-03",
     "Sync mode: Always On — syncs on interval regardless of user activity",
     "1. Set mode = Always On, interval = 5 min\n2. Actively use Revit for 6 min",
     "Sync triggers after 5 min even while user is active. Log: 'Starting sync'.",
     "High", "Both"),

    ("Sync Control", "Sync Settings", "SS-04",
     "Sync mode: When Paused — only fires during genuine idle",
     "1. Set mode = When Paused, idle timeout = 5 min\n2. Keep using Revit → wait 6 min",
     "No sync while user actively types/clicks. Sync fires only after 5 min of no input.",
     "High", "Both"),

    ("Sync Control", "Sync Settings", "SS-05",
     "Schedule window — sync blocked outside active window",
     "1. Set schedule 22:00–06:00\n2. Test at current time (outside window) → wait for idle",
     "No background sync triggered outside schedule window. Log: 'Outside schedule window'.",
     "Medium", "Both"),

    ("Sync Control", "Sync Settings", "SS-06",
     "Auto-exit: countdown dialog shown after idle threshold",
     "1. Enable Exit Revit on Idle, set to 3 min\n2. Leave machine idle 3+ min",
     "60-second countdown dialog appears: 'Revit will close in X seconds. Cancel?'",
     "High", "N/A"),

    ("Sync Control", "Sync Settings", "SS-07",
     "Auto-exit: cancel resets and re-arms for next idle period",
     "1. Countdown dialog appears → click Cancel",
     "Revit stays open. _exitRequested reset. After another idle period, dialog reappears.",
     "High", "N/A"),

    ("Sync Control", "Sync Settings", "SS-08",
     "Auto-exit: models synced before Revit closes",
     "1. Make changes to document\n2. Let countdown complete without cancelling",
     "Log shows sync for each modified doc before CloseMainWindow. No unsaved-changes dialog.",
     "High", "Both"),

    ("Sync Control", "Sync Settings", "SS-09",
     "Background relinquish fires only after idle >= RelinquishIntervalMinutes",
     "1. Enable relinquish, interval = 10 min\n2. Check out elements → idle 8 min (should NOT fire) → idle 12 min",
     "No relinquish at 8 min. Relinquish fires at 12 min. Log: 'Triggering relinquish'.",
     "High", "Both"),

    # ── 4b. Sync Queue (Activity Monitor) ───────────────────────────────────
    ("Sync Control", "Sync Queue", "SQ-01",
     "Sync Queue opens with last 15 sync records for current model",
     "1. Perform 5+ syncs\n2. ZeManage → Sync Queue",
     "Dialog shows historical records ordered latest first (max 15). Correct model filtered.",
     "High", "Both"),

    ("Sync Control", "Sync Queue", "SQ-02",
     "Live sync event prepended in real time",
     "1. Open Sync Queue dialog\n2. Perform a sync",
     "New row appears at top of list during sync (status: Syncing). Updates to Completed on finish.",
     "High", "Both"),

    ("Sync Control", "Sync Queue", "SQ-03",
     "Completed sync shows green badge and duration",
     "1. Perform a successful sync\n2. Check Sync Queue",
     "'Completed' green badge shown. Duration (e.g. '12s') displayed in Duration column.",
     "High", "Both"),

    ("Sync Control", "Sync Queue", "SQ-04",
     "Failed sync shows red badge",
     "1. Cause sync failure (disconnect network mid-sync)\n2. Check Sync Queue",
     "'Failed' red badge shown. Timestamp and user recorded.",
     "Medium", "Both"),

    ("Sync Control", "Sync Queue", "SQ-05",
     "Sync Traffic Control: manual sync conflict shows SyncConflictDialog",
     "1. Simulate another user actively syncing\n2. Click Synchronize With Central",
     "SyncConflictDialog shown. User can queue or cancel. NOT silently auto-joined.",
     "High", "Both"),

    ("Sync Control", "Sync Queue", "SQ-06",
     "Background sync conflict — silently queued without dialog",
     "1. Background sync triggers while another user is syncing",
     "No dialog shown. Background sync deferred. Log: '[BackgroundSyncEngine] Blocked by STC'.",
     "High", "Both"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 5. AI  (Ribbon Panel: "AI")
    #    Button: Ze AI
    # ═══════════════════════════════════════════════════════════════════════════

    ("AI", "Ze AI", "AI-01",
     "AI dialog opens from ribbon and initialises correctly",
     "1. ZeManage → Ze AI",
     "Chat dialog opens. Input field active. No error. Model context (current model name) shown.",
     "High", "N/A"),

    ("AI", "Ze AI", "AI-02",
     "Visibility issue diagnosis — specific solution provided",
     "1. Ask: 'Why is my element not visible in the current view?'\n2. Review response",
     "Response covers: View Range, Category Visibility, Workset visibility, Phase Filter, Crop Region, Discipline setting. Specific steps listed.",
     "High", "Both"),

    ("AI", "Ze AI", "AI-03",
     "Performance optimisation question uses correct knowledge section",
     "1. Ask: 'My model is very slow. How do I reduce file size?'\n2. Review response",
     "Response includes: purge unused, audit, large families, linked files, raster images. Matches Performance knowledge section.",
     "High", "Both"),

    ("AI", "Ze AI", "AI-04",
     "Workflow question — step-by-step answer provided",
     "1. Ask: 'How do I set up shared coordinates between two Revit models?'\n2. Review response",
     "Step-by-step workflow answer covering Acquire/Publish Coordinates, linked model coordination.",
     "Medium", "N/A"),

    ("AI", "Ze AI", "AI-05",
     "Best practices question — appropriate recommendations",
     "1. Ask: 'What are the best practices for naming Revit views?'\n2. Review response",
     "Response gives naming conventions, template usage, discipline prefixes. Matches BestPractices knowledge.",
     "Medium", "N/A"),

    ("AI", "Ze AI", "AI-06",
     "Model context injected into response",
     "1. Open a specific model\n2. Ask: 'How many warnings does this model have?'",
     "AI response references current model name. If warning count is available from metrics, it is cited.",
     "Medium", "Both"),

    ("AI", "Ze AI", "AI-07",
     "Conversation history maintained across messages",
     "1. Ask a question about Revit\n2. Ask follow-up: 'Can you elaborate?'\n3. Ask: 'Give me an example'",
     "AI maintains context across all 3 messages. Each follow-up builds on prior answer.",
     "Medium", "N/A"),

    ("AI", "Ze AI", "AI-08",
     "Off-topic question rejected gracefully",
     "1. Ask: 'What is the capital of France?'",
     "AI gives short, polite refusal scoped to Revit/BIM topics. No hallucination.",
     "Low", "N/A"),

    ("AI", "Ze AI", "AI-09",
     "AI unavailable / graceful fallback when not authenticated",
     "1. Log out\n2. Open Ze AI",
     "Dialog shows authentication-required message OR AI responds with limited non-model context.",
     "Medium", "N/A"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 6. ADDONS  (Ribbon Panel: "Addons" — only visible if BIManage.Addons.dll found)
    #    Buttons: NWC Export | Link Remapper
    # ═══════════════════════════════════════════════════════════════════════════

    ("Addons", "NWC Export", "NE-01",
     "Addons panel visible when BIManage.Addons.dll is present",
     "1. Confirm BIManage.Addons.dll exists in install folder\n2. Launch Revit",
     "Addons panel appears in ZeManage ribbon with NWC Export and Link Remapper buttons.",
     "High", "N/A"),

    ("Addons", "NWC Export", "NE-02",
     "NWC Export dialog opens and lists available 3D views",
     "1. Open a project model\n2. ZeManage → NWC Export",
     "Dialog opens. All 3D views in the project listed with checkboxes for selection.",
     "High", "Both"),

    ("Addons", "NWC Export", "NE-03",
     "Single 3D view exported to NWC successfully",
     "1. Select one 3D view → set output folder → Export",
     "NWC file created in output folder. File size > 0. Filename matches view name.",
     "High", "Both"),

    ("Addons", "NWC Export", "NE-04",
     "Batch export — multiple 3D views exported",
     "1. Select 3+ views → Export",
     "One NWC file created per selected view. All files valid. No partial failures.",
     "High", "Both"),

    ("Addons", "NWC Export", "NE-05",
     "Export of cloud model 3D views",
     "1. Open a BIM 360 / ACC model\n2. Export NWC",
     "Export succeeds from cloud model. NWC file created locally in output folder.",
     "Medium", "Cloud"),

    ("Addons", "Link Remapper", "LR-01",
     "Link Remapper dialog opens and lists all existing Revit links",
     "1. Open a project with Revit links\n2. ZeManage → Link Remapper",
     "Dialog opens. All linked models listed with current path, status (Loaded/Unloaded/Not Found).",
     "High", "Both"),

    ("Addons", "Link Remapper", "LR-02",
     "Remap local link to new local/network path",
     "1. Move linked RVT to new location\n2. Use Link Remapper to point to new path → Reload",
     "Link remapped. Model reloads from new path. Status changes to Loaded.",
     "High", "Local"),

    ("Addons", "Link Remapper", "LR-03",
     "Remap local link to cloud (BIM 360/ACC) path",
     "1. Link originally local → remap to cloud path via Link Remapper",
     "Link updated to cloud path. Revit reloads from cloud. Status = Loaded.",
     "High", "Cloud"),

    ("Addons", "Link Remapper", "LR-04",
     "Missing/broken link resolved via Link Remapper",
     "1. Open project with 'Not Found' link\n2. Use Link Remapper to locate the file",
     "Link resolved. File loads. Status changes from Not Found to Loaded.",
     "High", "Both"),

    ("Addons", "Link Remapper", "LR-05",
     "Remapping multiple links in one operation",
     "1. Project with 3+ links, 2 broken\n2. Remap both broken links → Reload All",
     "Both links remapped and reloaded in one operation. No remaining broken links.",
     "Medium", "Both"),

    # ═══════════════════════════════════════════════════════════════════════════
    # 7. SUPPORT  (Ribbon Panel: "Support")
    #    Buttons: Status (Device) | Account (Sign In) | Support
    # ═══════════════════════════════════════════════════════════════════════════

    # ── 7a. Status (Device Status) ──────────────────────────────────────────
    ("Support", "Status (Device)", "DS-01",
     "Status button shows RED when device is not registered",
     "1. Fresh install with no license\n2. Check ZeManage ribbon",
     "Status button is red. Tooltip: 'Red = not registered. Click to enter license key'.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-02",
     "Status button opens License Key dialog when not registered",
     "1. Click Status button (red state)",
     "License Key input dialog opens. User can enter license/registration key.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-03",
     "Valid license key registers device — button turns GREEN",
     "1. Enter valid license key → Submit",
     "Device registered. JWT token stored. Status button turns green. Log: 'Device validation successful'.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-04",
     "Invalid license key shows error — button remains RED",
     "1. Enter incorrect/expired license key → Submit",
     "Error message shown. No token stored. Button remains red.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-05",
     "Passive license — button turns YELLOW",
     "1. Device registered with passive license mode",
     "Status button is yellow. LicenseMode = Passive. Startup warning shown.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-06",
     "Status button (green) opens Device Status dialog",
     "1. Device registered (green button)\n2. Click Status",
     "Device Status dialog opens with: device ID, machine ID, company ID, registration date, license mode.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-07",
     "License module status shown (per-module on/off)",
     "1. Open Device Status dialog (or License Status dialog)\n2. Review module list",
     "Module status shown for: Protection, ActivityTracker, HealthMonitor, SyncControl, AI, Addons. ✓/✗ per module.",
     "High", "N/A"),

    ("Support", "Status (Device)", "DS-08",
     "License expiry date and seat count displayed",
     "1. Open Device Status / License Status dialog",
     "Expiry date shown. Seat usage (e.g. '3/10 seats') shown. Remaining days shown if near expiry.",
     "Medium", "N/A"),

    # ── 7b. Account (Sign In) ────────────────────────────────────────────────
    ("Support", "Account (Sign In)", "AC-01",
     "Sign In dialog opens and accepts credentials",
     "1. ZeManage → Account",
     "Login dialog opens. Email + password fields present. Sign In button active.",
     "High", "N/A"),

    ("Support", "Account (Sign In)", "AC-02",
     "Valid credentials — session created, ribbon commands enabled",
     "1. Enter valid email + password → Sign In",
     "JWT token obtained. Session created in DB. Ribbon commands (admin-only) become visible.",
     "High", "N/A"),

    ("Support", "Account (Sign In)", "AC-03",
     "Invalid credentials — error shown, no token stored",
     "1. Enter wrong password → Sign In",
     "Error message shown. No token stored. User remains logged out.",
     "High", "N/A"),

    ("Support", "Account (Sign In)", "AC-04",
     "Token auto-refresh on 401 (silent)",
     "1. Force expire access token (wait or manually)\n2. Perform any API action",
     "Token refreshed silently using refresh token. Action completes. No login prompt.",
     "High", "N/A"),

    ("Support", "Account (Sign In)", "AC-05",
     "Force Logout via SignalR — tokens cleared",
     "1. Admin sends ForceLogout from server console\n2. Observe client",
     "Client tokens cleared. Session marked Closed in DB. User prompted to sign in again.",
     "High", "N/A"),

    ("Support", "Account (Sign In)", "AC-06",
     "Already-signed-in state shows current user info",
     "1. Sign in successfully\n2. Click Account button again",
     "Account dialog shows current user: email, role, company, session ID.",
     "Medium", "N/A"),

    # ── 7c. Support ──────────────────────────────────────────────────────────
    ("Support", "Support Dialog", "SP-01",
     "Support dialog opens with contact info and version",
     "1. ZeManage → Support",
     "Dialog opens. Shows: ZeManage version, Revit version, support email/URL, company branding.",
     "High", "N/A"),

    ("Support", "Support Dialog", "SP-02",
     "Version number matches current build",
     "1. Open Support dialog → note version\n2. Compare with assembly version",
     "Version in dialog matches assembly version. No 'unknown' or '0.0.0' shown.",
     "Medium", "N/A"),

    ("Support", "Support Dialog", "SP-03",
     "Open Log button opens latest log file",
     "1. Click 'Open Log' in Support dialog",
     "Latest BIManageRevit_*.log file opens in default text editor. File is not empty.",
     "Medium", "N/A"),
]

# ── EXCEL BUILDER ─────────────────────────────────────────────────────────────
def write_checklist():
    wb = openpyxl.Workbook()

    # ── Instructions ─────────────────────────────────────────────────────────
    cover = wb.active
    cover.title = "Instructions"
    cover.sheet_view.showGridLines = False
    cover.column_dimensions["A"].width = 2
    cover.column_dimensions["B"].width = 90

    cover["B2"] = "ZeManage Revit Add-in — Test Checklist  v2.0"
    cover["B2"].font = Font(name="Calibri", bold=True, size=18, color="1E293B")
    cover["B3"] = "Mapped to ZeManage ribbon panels and buttons  |  Generated: 2026-03-24"
    cover["B3"].font = Font(name="Calibri", size=11, color="64748B")

    lines = [
        "", "HOW TO USE",
        "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
        "1. Work through the 'Test Checklist' tab row by row.",
        "2. Set Status: Pass / Fail / Blocked / Skip",
        "3. Fill Actual Result / Notes for any Fail or Blocked.",
        "4. Enter Tester name, Date Tested, Build/Version per row.",
        "5. Model Type column: 'Local', 'Cloud', or 'Both' — run the test against the appropriate model.",
        "",
        "STATUS CODES",
        "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
        "  Pass     — Test executed, result matches expected.",
        "  Fail     — Result does not match expected. Document in Notes.",
        "  Blocked  — Cannot execute (environment/dependency issue).",
        "  Skip     — Not in scope for this build/sprint.",
        "",
        "RIBBON PANELS COVERED",
        "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
        "  1. Protection       — Pin Protection | Command Protection | Event Restriction | Rule Management",
        "  2. Activity Tracker — Session Info | Crash Register | Model Activities",
        "  3. Health Monitor   — Model Health (Dashboard + scoring) | Analyze Model (expensive metrics)",
        "                        Metrics types: Sync/Save (auto fast) | Periodic (daily medium) | Manual (expensive)",
        "  4. Sync Control     — Sync Settings | Sync Queue",
        "  5. AI               — Ze AI (OpenAI-powered, model context aware)",
        "  6. Addons           — NWC Export | Link Remapper",
        "  7. Support          — Status (device/license) | Account (sign in) | Support",
    ]
    for i, line in enumerate(lines, start=4):
        c = cover.cell(row=i, column=2, value=line)
        if line in ("HOW TO USE", "STATUS CODES", "RIBBON PANELS COVERED"):
            c.font = Font(name="Calibri", bold=True, size=11, color="1E293B")
        elif line.startswith("━"):
            c.font = Font(name="Calibri", size=9, color="CBD5E1")
        else:
            c.font = Font(name="Calibri", size=10, color="374151")
        cover.row_dimensions[i].height = 16

    # ── Checklist sheet ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Test Checklist")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "E3"

    for col_letter, _, width in COLUMNS:
        ws.column_dimensions[col_letter].width = width

    # Row 1 — Title
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A1:N1")
    t = ws["A1"]
    t.value = "ZeManage Revit Add-in — Test Checklist v2.0  |  ZeManage Ribbon Panels"
    t.fill = fill(HEADER_BG)
    t.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    t.alignment = Alignment(horizontal="center", vertical="center")

    # Row 2 — Headers
    ws.row_dimensions[2].height = 22
    for idx, (_, col_name, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=2, column=idx)
        c.value = col_name
        c.fill = fill("334155")
        c.font = Font(name="Calibri", bold=True, size=9, color="F8FAFC")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = med()

    # Dropdowns
    for col_letter, formula in [("J", '"Pass,Fail,Blocked,Skip"'), ("I", '"Local,Cloud,Both,N/A"')]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        dv.sqref = f"{col_letter}3:{col_letter}1000"
        ws.add_data_validation(dv)

    row_num = 3
    seq = 1
    current_module = None

    for (module, submodule, tc_id, desc, steps, expected, priority, model_type) in TESTS:
        fg, bg = MOD_COLORS.get(module, ("374151", "F1F5F9"))

        # Module separator row
        if module != current_module:
            current_module = module
            ws.row_dimensions[row_num].height = 19
            ws.merge_cells(f"A{row_num}:N{row_num}")
            sep = ws.cell(row=row_num, column=1)
            sep.value = f"  {module.upper()}"
            sep.fill = fill(fg)
            sep.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
            sep.alignment = Alignment(vertical="center")
            row_num += 1

        ws.row_dimensions[row_num].height = 65

        values = [seq, module, submodule, tc_id, desc, steps, expected,
                  priority, model_type, "", "", "", "", ""]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_num, column=col_idx)
            c.value = val
            c.border = thin()

            # Background
            if col_idx == 10:   # Status
                c.fill = fill("FFFFFF")
            elif col_idx == 9:  # Model Type
                mt_fills = {"Local": "F0FDF4", "Cloud": "EFF6FF", "Both": "FFFBEB", "N/A": "F8FAFC"}
                c.fill = fill(mt_fills.get(val, "FFFFFF"))
            else:
                c.fill = fill(bg if seq % 2 == 1 else "FFFFFF")

            # Alignment
            center_cols = {1, 4, 8, 9, 10, 12, 13, 14}
            c.alignment = Alignment(
                horizontal="center" if col_idx in center_cols else "left",
                vertical="top",
                wrap_text=True
            )

            # Font
            if col_idx == 4:   # TC ID
                c.font = Font(name="Calibri", bold=True, size=9, color=fg)
            elif col_idx == 8: # Priority
                p = {"High": ("FEF2F2", "991B1B"), "Medium": ("FEF3C7", "92400E"), "Low": ("F0FDF4", "166534")}
                pb, pf = p.get(val, ("F1F5F9", "475569"))
                c.fill = fill(pb)
                c.font = Font(name="Calibri", bold=True, size=9, color=pf)
            elif col_idx == 9: # Model Type
                mt_fg = {"Local": "166534", "Cloud": "1D4ED8", "Both": "92400E", "N/A": "475569"}
                c.font = Font(name="Calibri", bold=True, size=9, color=mt_fg.get(val, "475569"))
            else:
                c.font = Font(name="Calibri", size=9, color="1E293B")

        seq += 1
        row_num += 1

    # ── Summary sheet ─────────────────────────────────────────────────────────
    summ = wb.create_sheet("Summary")
    summ.sheet_view.showGridLines = False
    for col, w in [("A",2),("B",24),("C",11),("D",11),("E",11),("F",11),("G",11),("H",20)]:
        summ.column_dimensions[col].width = w

    summ.row_dimensions[1].height = 28
    summ.merge_cells("B1:H1")
    st = summ["B1"]
    st.value = "Test Execution Summary"
    st.fill = fill(HEADER_BG)
    st.font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    st.alignment = Alignment(horizontal="center", vertical="center")

    for i, h in enumerate(["Module","Total","Pass","Fail","Blocked","Skip","Notes"], start=2):
        c = summ.cell(row=2, column=i)
        c.value = h
        c.fill = fill("334155")
        c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    summ.row_dimensions[2].height = 20

    module_counts = Counter(t[0] for t in TESTS)
    sr = 3
    for mod, (fg, bg) in MOD_COLORS.items():
        total = module_counts.get(mod, 0)
        summ.row_dimensions[sr].height = 18
        for ci, val in enumerate([mod, total,"","","","",""], start=2):
            c = summ.cell(row=sr, column=ci)
            c.value = val
            c.fill = fill(bg)
            c.font = Font(name="Calibri", bold=(ci==2), size=10, color="1E293B")
            c.alignment = Alignment(horizontal="center" if ci>2 else "left", vertical="center")
            c.border = thin()
        sr += 1

    summ.row_dimensions[sr].height = 20
    for ci, val in enumerate(["TOTAL", len(TESTS),"","","","",""], start=2):
        c = summ.cell(row=sr, column=ci)
        c.value = val
        c.fill = fill(HEADER_BG)
        c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        c.alignment = Alignment(horizontal="center" if ci>2 else "left", vertical="center")

    out = r"C:\Users\Admin\Documents\GitHub_JC\BIManageRevit\Tools\ZeManage_Test_Checklist.xlsx"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Total: {len(TESTS)} test cases")
    for mod, cnt in module_counts.items():
        print(f"  {mod}: {cnt}")

if __name__ == "__main__":
    write_checklist()
